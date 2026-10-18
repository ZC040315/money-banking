"""把 pptx / docx 的正文按页/段落抽成纯文本，便于逐章对齐。

用法：
    python scripts/extract_office.py <文件或目录> [-o 输出目录]

输出：<输出目录>/<原文件名>.txt，默认写到同级 _extracted/ 下。
表格会以 `| a | b |` 形式保留；图片只记录数量（图片请另发，我直接看图）。
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

# Windows 控制台默认 GBK，先把标准输出切到 UTF-8，避免中文与符号报错
try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:  # pragma: no cover - 老版本 Python 没有 reconfigure
    pass


def extract_pptx(path: Path) -> str:
    from pptx import Presentation

    prs = Presentation(str(path))
    lines: list[str] = [f"# {path.name}", f"（共 {len(prs.slides)} 页）", ""]
    for index, slide in enumerate(prs.slides, start=1):
        lines.append(f"## 第 {index} 页")
        texts: list[str] = []
        tables: list[str] = []
        pictures = 0
        for shape in slide.shapes:
            if shape.shape_type == 13 or shape.__class__.__name__ == "Picture":
                pictures += 1
                continue
            if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip():
                texts.append(shape.text_frame.text.strip())
            if getattr(shape, "has_table", False):
                rows = []
                for row in shape.table.rows:
                    rows.append("| " + " | ".join(cell.text.strip() for cell in row.cells) + " |")
                tables.append("\n".join(rows))
        if texts:
            lines.append("\n".join(texts))
        if tables:
            lines.append("")
            lines.extend(tables)
        if pictures:
            lines.append(f"（本页含 {pictures} 张图片/图表）")
        lines.append("")
    return "\n".join(lines)


def extract_docx(path: Path) -> str:
    from docx import Document

    doc = Document(str(path))
    lines: list[str] = [f"# {path.name}", ""]
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style = (paragraph.style.name or "").lower()
        if style.startswith("heading"):
            level = "".join(ch for ch in style if ch.isdigit()) or "2"
            lines.append(f"{'#' * min(int(level) + 1, 6)} {text}")
        else:
            lines.append(text)
    for table_index, table in enumerate(doc.tables, start=1):
        lines.append("")
        lines.append(f"表 {table_index}：")
        for row in table.rows:
            lines.append("| " + " | ".join(cell.text.strip() for cell in row.cells) + " |")
    return "\n".join(lines)


def convert(path: Path, out_dir: Path) -> None:
    suffix = path.suffix.lower()
    if suffix == ".pptx":
        content = extract_pptx(path)
    elif suffix == ".docx":
        content = extract_docx(path)
    else:
        print(f"skip (unsupported): {path.name}")
        return
    out_dir.mkdir(parents=True, exist_ok=True)
    target = out_dir / f"{path.stem}.txt"
    target.write_text(content, encoding="utf-8")
    print(f"[ok] {path.name} -> {target.name}  ({len(content)} chars)")


def main() -> int:
    parser = argparse.ArgumentParser(description="抽取 pptx/docx 文本")
    parser.add_argument("source", help="文件或目录")
    parser.add_argument("-o", "--out", default=None, help="输出目录，默认 <源目录>/_extracted")
    args = parser.parse_args()

    source = Path(args.source)
    if not source.exists():
        print(f"路径不存在：{source}")
        return 1

    out_dir = Path(args.out) if args.out else (source if source.is_dir() else source.parent) / "_extracted"
    files = [source] if source.is_file() else sorted(
        p for p in source.rglob("*") if p.suffix.lower() in {".pptx", ".docx"}
    )
    if not files:
        print("no .pptx / .docx found. For PDF, send screenshots instead.")
        return 1
    for path in files:
        convert(path, out_dir)
    print(f"\ndone: {len(files)} file(s) -> {out_dir}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
